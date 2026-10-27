"""
genera_lettere.py — lettere di fine collaudo (Trasmissione copia verbale di collaudo · Richiesta estensione in
esercizio) compilate dai template Word `Template/Template - ADMXXX-XX-XX <Lettera>.docx`, con i dati del rilascio.
Allineato alla versione web (42_lettere.js / generaLettere, spec docs/superpowers/specs/2026-10-21-leviathadm-web-lettere-design.md).

Uso da codice:
    from genera_lettere import dati_lettere, genera_lettere
    dati = dati_lettere("ADM956-01-0C")          # precompilazione: scheda, resoconto di collaudo, Verbale, Anagrafica
    dati["adm"].append("Rossi Mario")              # tutto modificabile prima di generare
    genera_lettere("ADM956-01-0C", ["trasmissione", "estensione"], dati, sovrascrivi=False)
    # → {"esiti": [{"tipo", "path", "stato": creata|esistente|sovrascritta, "svuotati": [...]}], "avvisi": [...]}

Regole (CLAUDE.md "Lettere di fine collaudo"): solo i run segnaposto vengono sostituiti (tab, rientri, firma del
Dirigente intatti); una lettera già presente (qualunque `*<Lettera>.docx` nella root, come build_da_firmare) non si
tocca senza `sovrascrivi`; con `sovrascrivi` si riscrive e si svuota `da mandare in firma/<Lettera>/` (prepara_firma
rifà i PDF solo se la cartella è vuota), con avviso se la lettera è già segnata in firma. I PDF restano a prepara_firma.
"""
from __future__ import annotations

import copy
import io
import re
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from mail_utils import parse_frontmatter, participants_from_resoconto, read_release, team_from_release  # noqa: E402
from parse_presenze import find_presenze_csv, parse_presenze_csv  # noqa: E402
from update_anagrafica import anagrafica_info  # noqa: E402

VAULT = Path(__file__).parent.parent
LAVORI = VAULT / "PO 2026"
TEMPLATE_DIR = VAULT / "Template"
ANAGRAFICA = VAULT / "Anagrafica"

LETTERE = {"trasmissione": "Trasmissione copia verbale di collaudo", "estensione": "Richiesta estensione in esercizio"}
FIELD_BASE = {"trasmissione": "trasmissione_verbale", "estensione": "estensione_esercizio"}
ESITI = ("POSITIVO", "con riserva", "NEGATIVO")
_MAPPA_ESITO = {"positivo": "POSITIVO", "con riserva": "con riserva", "negativo": "NEGATIVO"}
PARTICELLE = set("a e di del della delle dei degli ed per in al alla ai agli su sul".split())
PARTICELLE_COGNOME = set("de di del della dello dei degli da dal dalla dalle lo la le li van von mc mac san santa".split())
AVVISO_DENOMINAZIONE = "verifica la denominazione nell'intestazione"
AVVISO_RB = "RB senza direzione/ufficio in Anagrafica: compila a mano (e la scheda della persona)"
AVVISO_EFFETTIVA = "non compariranno in Da firmare finché il collaudo non è registrato come svolto"
_SEGNAPOSTO = ("[Nome]", "[Cognome]", "[Descrizione", "[Direzione", "[Ufficio", "[Area")


# ── nomi e denominazioni ──────────────────────────────────────────────────────

def nome_file(codice: str, tipo: str) -> str:
    """Convenzione dei file nel vault: Trasmissione senza trattino, Richiesta estensione col trattino."""
    return f"{codice} {LETTERE[tipo]}.docx" if tipo == "trasmissione" else f"{codice} - {LETTERE[tipo]}.docx"


def cognome_nome(nome: str) -> tuple[str, str]:
    """"Cognome Nome" dell'Anagrafica → (cognome, nome). "Cognome | Nome" forza la divisione; particella iniziale
    (De, Di, Lo, Della…) + cognome quando restano almeno tre parole; suffissi tra parentesi rimossi."""
    n = " ".join(re.sub(r"\s*\([^)]*\)", "", nome or "").split())
    if "|" in n:
        c, _, r = n.partition("|")
        return c.strip(), r.strip()
    parti = n.split(" ") if n else []
    if not parti:
        return "", ""
    k = 2 if len(parti) >= 3 and parti[0].lower() in PARTICELLE_COGNOME else 1
    return " ".join(parti[:k]), " ".join(parti[k:])


def ufficio_da_organigramma(denominazione: str) -> str:
    """"GIOCO A DISTANZA E SCOMMESSE" → "Ufficio Gioco a Distanza e Scommesse" (convenzione: da verificare nel Word)."""
    parole = [p for p in (denominazione or "").strip().split() if p]
    if not parole:
        return ""
    out = []
    for i, p in enumerate(parole):
        low = p.lower()
        out.append(low if (i > 0 and low in PARTICELLE) else low[:1].upper() + low[1:])
    return "Ufficio " + " ".join(out)


def direzione_da_sigla(sigla: str) -> str:
    """nome_ufficiale di Anagrafica/Strutture/<sigla>.md, altrimenti la sigla."""
    s = (sigla or "").strip()
    if not s:
        return ""
    nodo = ANAGRAFICA / "Strutture" / f"{s}.md"
    nome = parse_frontmatter(nodo).get("nome_ufficiale", "").strip() if nodo.exists() else ""
    return nome or s


def esito_da_resoconto(testo_md: str) -> str:
    m = re.search(r"esito \*\*(POSITIVO|con riserva|NEGATIVO)\*\*", testo_md or "")
    return m.group(1) if m else ""


def _esito_da_verbale(path: Path) -> str:
    fm = parse_frontmatter(path)
    v = (fm.get("esito_collaudo", "") or "").strip().lower()
    if v in _MAPPA_ESITO:
        return _MAPPA_ESITO[v]
    testo = path.read_text(encoding="utf-8", errors="replace")
    sez = testo.split("## Esito complessivo", 1)
    if len(sez) == 2:
        for riga, esito in (("- [x] Superato con riserva", "con riserva"), ("- [x] Non superato", "NEGATIVO"), ("- [x] Superato", "POSITIVO")):
            if riga.lower() in sez[1].lower():
                return esito
    return ""


# ── dati del rilascio ─────────────────────────────────────────────────────────

def _data_nel_nome(p: Path) -> str:
    m = re.search(r" - (\d{8}) ", p.name)
    return m.group(1) if m else ""


def dati_lettere(codice: str, oggi: date | None = None) -> dict:
    """Precompilazione (come LV.lettere.datiLettereConCsv): tabella PO dalla scheda; rappresentanti dall'ultimo resoconto
    di collaudo (root o inviati/), altrimenti dal CSV presenze della data del collaudo, altrimenti dal team della scheda;
    esito dal Verbale, poi dal resoconto, poi POSITIVO; Direzione/Ufficio del RB dall'Anagrafica."""
    oggi = oggi or date.today()
    fm, testo_scheda = read_release(codice)
    release_dir = LAVORI / codice
    team = team_from_release(testo_scheda)
    avvisi: list[str] = []
    collaudi = sorted([p for p in list(release_dir.glob(f"{codice} - * Resoconto Collaudo*.md"))
                       + list((release_dir / "inviati").glob(f"{codice} - * Resoconto Collaudo*.md"))], key=_data_nel_nome)
    ultimo = collaudi[-1] if collaudi else None
    adm: list[str] = []
    sogei: list[str] = []
    fonte = ""
    if ultimo is not None:
        a, s = participants_from_resoconto(ultimo)
        if a or s:
            adm, sogei, fonte = a, s, "resoconto"
    if not fonte:
        effettiva = (fm.get("data_collaudo_effettiva", "") or "").strip()
        csv = find_presenze_csv(release_dir, effettiva) if effettiva else None
        if csv is not None:
            pr = parse_presenze_csv(csv)
            if pr.get("adm") or pr.get("sogei"):
                adm, sogei, fonte = [p["nome"] for p in pr["adm"]], [p["nome"] for p in pr["sogei"]], "csv"
    if not fonte:
        adm = [n for r in ("PM", "RA", "RP", "RB") for n in team.get(r, [])]
        sogei = list(team.get("SOGEI", []))
        ref = (fm.get("referente_sogei", "") or "").strip()
        if ref and ref not in sogei:
            sogei.insert(0, ref)
        fonte = "team" if (adm or sogei) else ""
    pulisci = lambda n: " ".join(cognome_nome(n)).strip()   # noqa: E731
    adm = [pulisci(n) for n in adm if pulisci(n) and not n.strip().startswith("[")]
    sogei = [pulisci(n) for n in sogei if pulisci(n) and not n.strip().startswith("[")]
    verbali = sorted(release_dir.glob(f"{codice} - * Verbale di Collaudo.md"), key=_data_nel_nome)
    esito = _esito_da_verbale(verbali[-1]) if verbali else ""
    if not esito and ultimo is not None:
        esito = esito_da_resoconto(ultimo.read_text(encoding="utf-8", errors="replace"))
    esito = esito or "POSITIVO"
    rb = team.get("RB", [""])[0] if team.get("RB") else ""
    direzione = ufficio = ""
    if rb:
        info = anagrafica_info(rb)
        direzione = direzione_da_sigla(info.get("direzione", ""))
        ufficio = ufficio_da_organigramma(info.get("ufficio", ""))
    if not direzione or not ufficio:
        avvisi.append(AVVISO_RB)
    avvisi.append(AVVISO_DENOMINAZIONE)
    eff = (fm.get("data_collaudo_effettiva", "") or "").strip()
    try:
        eff_d = date.fromisoformat(eff) if eff else None
    except ValueError:
        eff_d = None
    if eff_d is None or eff_d > oggi:
        avvisi.append(AVVISO_EFFETTIVA)
    return {"progetto": fm.get("codice_progetto", "").strip(), "descrizione_progetto": fm.get("descrizione_progetto", "").strip(),
            "sottoprogetto": fm.get("codice_sottoprogetto", "").strip(),
            "rilascio": f"{codice} – {fm.get('descrizione_rilascio', '').strip()}",
            "adm": adm, "sogei": sogei, "fonte_nomi": fonte, "esito": esito, "direzione": direzione, "ufficio": ufficio,
            "area_staff": "Area di Staff", "rb": rb, "avvisi": avvisi}


# ── riempimento dei template (python-docx, come lo script del 2026-10-20) ──────

def _set_cell(cell, text: str) -> None:
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)
    for extra in cell.paragraphs[1:]:
        for r in extra.runs:
            r.text = ""


def _tabella_po(doc, dati: dict) -> None:
    t = doc.tables[0]
    for cell, v in zip(t.rows[1].cells, [dati["progetto"], dati["descrizione_progetto"], dati["sottoprogetto"], dati["rilascio"]]):
        _set_cell(cell, str(v))


def _tabella_persone(t, nomi: list[str]) -> None:
    righe = t.rows[1:]
    if not righe:
        raise ValueError("template Word non conforme: tabella dei rappresentanti senza righe segnaposto")
    while len(righe) < len(nomi):
        t._tbl.append(copy.deepcopy(righe[-1]._tr))
        righe = t.rows[1:]
    for row, nome in zip(righe, nomi):
        cognome, proprio = cognome_nome(nome)
        _set_cell(row.cells[0], cognome)
        _set_cell(row.cells[1], proprio)
    for row in righe[len(nomi):]:
        row._tr.getparent().remove(row._tr)


def _sostituisci_run(doc, mappa: dict[str, str]) -> None:
    for p in doc.paragraphs:
        for r in p.runs:
            for old, new in mappa.items():
                if old in r.text:
                    r.text = r.text.replace(old, new)


def _residui(doc) -> list[str]:
    testi = [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return [t for t in testi if any(s in t for s in _SEGNAPOSTO) and not t.startswith("Prot.:")]


def _bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _controlla(doc) -> None:
    res = _residui(doc)
    if res:
        raise ValueError("template Word non conforme: segnaposto non sostituito: " + " | ".join(res)[:200])


def riempi_trasmissione(template: Path, dati: dict) -> bytes:
    import docx
    d = docx.Document(str(template))
    _tabella_po(d, dati)
    _tabella_persone(d.tables[1], list(dati.get("adm", [])))
    _tabella_persone(d.tables[2], list(dati.get("sogei", [])))
    for p in d.paragraphs:
        if "esito POSITIVO" in p.text:
            for r in p.runs:
                if "POSITIVO" in r.text:
                    r.text = r.text.replace("POSITIVO", dati.get("esito") or "POSITIVO")
                    break
            break
    else:
        raise ValueError('template Word non conforme: paragrafo "Il collaudo ha avuto esito POSITIVO" assente')
    _controlla(d)
    return _bytes(d)


def riempi_estensione(template: Path, dati: dict) -> bytes:
    import docx
    d = docx.Document(str(template))
    _tabella_po(d, dati)
    _sostituisci_run(d, {"[Direzione di business]": dati.get("direzione", ""), "[Ufficio del Responsabile di Business]": dati.get("ufficio", ""),
                         "[Area di Staff]": dati.get("area_staff") or "Area di Staff"})
    _controlla(d)
    return _bytes(d)


# ── scrittura nel vault ───────────────────────────────────────────────────────

def lettera_esistente(release_dir: Path, label: str) -> Path | None:
    """Stesso glob di build_da_firmare: qualunque *<Lettera>.docx nella root (maiuscole indifferenti), anche con l'altra forma del nome."""
    suff = f"{label}.docx".lower()
    cand = sorted(p for p in release_dir.iterdir() if p.is_file() and p.name.lower().endswith(suff))
    return cand[0] if cand else None


def genera_lettere(codice: str, lettere: list[str], dati: dict | None = None, *, sovrascrivi: bool = False,
                   out_dir: Path | None = None) -> dict:
    """Scrive le lettere scelte nella root del rilascio. Ritorna {"esiti": [{tipo, path, stato, svuotati}], "avvisi": [...]}."""
    dati = dati or dati_lettere(codice)
    fm, _ = read_release(codice)
    release_dir = out_dir or (LAVORI / codice)
    if "trasmissione" in lettere and not (dati.get("adm") and dati.get("sogei")):
        raise ValueError("almeno un nome ADM e uno SOGEI per la Trasmissione copia verbale")
    if "estensione" in lettere and not ((dati.get("direzione") or "").strip() and (dati.get("ufficio") or "").strip()):
        raise ValueError("Direzione e Ufficio del referente business sono obbligatori per la Richiesta estensione in esercizio (riga e, p.c.)")
    pronti = []
    for tipo in lettere:
        label = LETTERE[tipo]
        tpl = TEMPLATE_DIR / f"Template - ADMXXX-XX-XX {label}.docx"
        if not tpl.exists():
            raise ValueError(f"template assente nel vault: {tpl.relative_to(VAULT)}")
        pronti.append((tipo, label, riempi_trasmissione(tpl, dati) if tipo == "trasmissione" else riempi_estensione(tpl, dati)))
    esiti, avvisi = [], []
    for tipo, label, contenuto in pronti:
        esistente = lettera_esistente(release_dir, label)
        if esistente is not None and not sovrascrivi:
            esiti.append({"tipo": tipo, "path": esistente, "stato": "esistente", "svuotati": []})
            continue
        svuotati: list[Path] = []
        if esistente is not None and (fm.get(f"{FIELD_BASE[tipo]}_in_firma", "") or "").strip().lower() == "true":
            avvisi.append(f"{label}: lettera già in firma — riportala in Da mandare (campo {FIELD_BASE[tipo]}_in_firma nella scheda) perché il sync del master rifaccia i PDF")
        path = esistente or (release_dir / nome_file(codice, tipo))
        path.write_bytes(contenuto)
        if esistente is not None:   # dopo la lettera: prepara_firma rifà i PDF solo se la cartella è vuota
            cartella = release_dir / "da mandare in firma" / label
            if cartella.is_dir():
                for f in sorted(p for p in cartella.rglob("*") if p.is_file()):
                    f.unlink()
                    svuotati.append(f)
        esiti.append({"tipo": tipo, "path": path, "stato": "sovrascritta" if esistente is not None else "creata", "svuotati": svuotati})
    return {"esiti": esiti, "avvisi": avvisi}


if __name__ == "__main__":
    print(__doc__)
